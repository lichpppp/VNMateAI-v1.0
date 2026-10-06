"""
mateai/application/enterprise/erp_import.py
===========================================
Nhập cấu trúc tổ chức ERP từ Excel + file mẫu (Phase 47). Chuyển từ
`interfaces/http/api_erp.py` (Supervisor Phase 10, §198). Hàm ĐỒNG BỘ (pandas,
openpyxl, SQLite) — tầng HTTP gọi qua `run_blocking`, không chạy trên event loop.

Chỉ nhận `.xlsx`: trước đây nhận cả `.csv` / `.xls` nhưng đọc bằng openpyxl (chỉ đọc
được .xlsx), nên hai định dạng đó luôn hỏng với lỗi khó hiểu.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any, Dict, List, Tuple

logger = logging.getLogger(__name__)

#: Phase 73: file mẫu CHỈ có hàng tiêu đề — không dòng dữ liệu bịa (nhân sự / IP giả
#: nạp vào sẽ thành dữ liệu "thật" trong hệ thống).
TEMPLATE_COLUMNS: Dict[str, List[str]] = {
    "PhongBan": ["TenPhongBan", "MoTa"],
    "NhanVien": ["PhongBan", "HoTen", "ChucVu", "Email", "SoDienThoai"],
    "MayTinh": ["PhongBan", "TenMay", "DiaChiIP", "LoaiThietBi", "NguoiSuDung"],
    "CongViec": ["PhongBan", "TieuDe", "NguoiPhuTrach", "TrangThai", "HanChot"],
    "SoSach": ["PhongBan", "TenTaiLieu", "DuongDan", "NgayTao"],
}

#: Tên sheet chấp nhận (không phân biệt hoa thường) cho từng loại dữ liệu.
SHEET_ALIASES: Dict[str, List[str]] = {
    "departments_data": ["PhongBan", "Departments", "Phòng Ban", "Phong_Ban"],
    "employees_data": ["NhanVien", "Employees", "Nhân Viên", "Nhan_Vien"],
    "devices_data": ["MayTinh", "Devices", "Máy Tính", "May_Tinh", "ThietBi"],
    "tasks_data": ["CongViec", "Tasks", "Công Việc", "Cong_Viec"],
    "records_data": ["SoSach", "Records", "Sổ Sách", "So_Sach", "HoSo", "Hồ Sơ"],
}

ACCEPTED_EXT = (".xlsx",)


class ImportRejected(Exception):
    """File không nhập được (sai định dạng / rỗng / thiếu sheet) — chưa ghi gì."""


def build_template() -> Tuple[bytes, str]:
    """(nội dung .xlsx, tên file) — mỗi sheet chỉ có hàng tiêu đề, đã định dạng."""
    import pandas as pd
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    output = io.BytesIO()
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    side = Side(style="thin", color="CCCCCC")
    border_thin = Border(left=side, right=side, top=side, bottom=side)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, columns in TEMPLATE_COLUMNS.items():
            pd.DataFrame({col: [] for col in columns}).to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            ws.views.sheetView[0].showGridLines = True
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 28
            for col in ws.columns:
                max_len = 0
                for idx, cell in enumerate(col):
                    cell.border = border_thin
                    if idx > 0:
                        cell.font = Font(name="Arial", size=10)
                        cell.alignment = Alignment(vertical="center")
                    max_len = max(max_len, len(str(cell.value or "")))
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 5, 18)
    return output.getvalue(), f"Mau_Nhap_Lieu_ERP_VNMATEAI_{datetime.now().strftime('%Y%m%d')}.xlsx"


def parse_workbook(contents: bytes) -> Dict[str, List[Dict[str, Any]]]:
    """Đọc các sheet đã biết thành danh sách dòng (ô trống -> "")."""
    import pandas as pd

    xls = pd.ExcelFile(io.BytesIO(contents), engine="openpyxl")
    names = xls.sheet_names

    def read(candidates: List[str]) -> List[Dict[str, Any]]:
        for c in candidates:
            for actual in names:
                if actual.strip().lower() == c.strip().lower():
                    return pd.read_excel(xls, sheet_name=actual).fillna("").to_dict(orient="records")
        return []

    return {key: read(aliases) for key, aliases in SHEET_ALIASES.items()}


def import_workbook(actor: str, filename: str, contents: bytes) -> Dict[str, Any]:
    """Nhập trong MỘT transaction (erp_db rollback toàn bộ nếu một dòng lỗi).

    Ném `ImportRejected` (file không dùng được) hoặc `ValueError` (dòng dữ liệu sai,
    thông điệp chỉ rõ dòng — trả nguyên cho người dùng).
    """
    from mateai.infrastructure.database.erp_database import erp_db

    if not (filename or "").lower().endswith(ACCEPTED_EXT):
        raise ImportRejected("Chỉ nhận file Excel .xlsx (tải file mẫu để xem đúng cấu trúc). "
                             "File .xls / .csv cần mở bằng Excel rồi lưu lại dạng .xlsx.")
    if not contents:
        raise ImportRejected("Tệp tải lên rỗng.")
    try:
        sheets = parse_workbook(contents)
    except Exception as exc:  # noqa: BLE001 — file hỏng / không phải xlsx thật
        raise ImportRejected(f"Không đọc được file Excel: {exc}") from exc
    if not any(sheets[k] for k in ("departments_data", "employees_data", "devices_data", "tasks_data")):
        raise ImportRejected("Không tìm thấy các Sheet dữ liệu chuẩn ('PhongBan', 'NhanVien', 'MayTinh', "
                             "'CongViec'). Vui lòng tải file mẫu để kiểm tra.")

    result = erp_db.bulk_import(**sheets)
    logger.info("Người dùng '%s' đã import dữ liệu ERP thành công: %s", actor, result.get("stats"))
    _audit(actor, "erp_import", {"file": filename, "stats": result.get("stats")})
    return result


def _audit(actor: str, action: str, details: Dict[str, Any]) -> None:
    try:
        from mateai.application.security.safety_guard import security_engine
        security_engine.log_audit(actor, action, "ERP", "SUCCESS", details)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Không ghi được audit %s: %s", action, exc)
